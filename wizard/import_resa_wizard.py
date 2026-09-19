import base64
import io
import logging
import re
from datetime import datetime
from openpyxl import load_workbook
from odoo import models, fields, api, _
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)

BATCH_SIZE = 2000

# === NOUVEAUX TARIFS (€/kWh HTVA) ===
CV_RESTITUTION_PRICE = 0.026
ACCISES_2026_PRICE = 0.046


class ImportResaWizard(models.TransientModel):
    _name = 'import.resa.wizard'
    _description = 'Assistant Import RESA'

    file_name = fields.Char(string='Nom du fichier')
    file_content = fields.Binary(string='Fichier Excel', required=True)
    community_id = fields.Many2one('energy.community',
                                   string='Opération de partage', required=True)
    month = fields.Char(
        string='Mois (YYYY-MM)',
        default=lambda self: datetime.now().strftime('%Y-%m')
    )

    generate_invoices = fields.Boolean(
        string='Générer les factures automatiquement',
        default=False,
    )
    invoice_date = fields.Date(
        string='Date de facture',
        default=fields.Date.today,
    )
    payment_term_id = fields.Many2one(
        'account.payment.term',
        string='Condition de paiement',
    )

    new_participants = fields.Text(string='Nouveaux participants', readonly=True)
    existing_participants = fields.Text(string='Participants existants', readonly=True)
    has_new_participants = fields.Boolean(string='A des nouveaux participants', default=False)
    confirm_import = fields.Boolean(
        string='Je confirme l\'ajout des nouveaux participants', default=False)
    import_enabled = fields.Boolean(string='Import activé', default=False)

    import_stats = fields.Text(string='Statistiques d\'import', readonly=True)
    invoices_created = fields.Integer(string='Factures créées', default=0)
    invoices_skipped = fields.Integer(string='Factures déjà existantes', default=0)

    # ==========================================
    # ANALYSE
    # ==========================================

    @api.onchange('file_content')
    def _onchange_file_content(self):
        self.new_participants = False
        self.existing_participants = False
        self.has_new_participants = False
        self.confirm_import = False
        self.import_enabled = False
        self.import_stats = False

        if not self.file_content:
            return

        try:
            self._detect_month()
            excel_data = base64.b64decode(self.file_content)
            workbook = load_workbook(io.BytesIO(excel_data), data_only=True)
            ean_list = self._extract_eans_from_sheets(workbook)

            if not ean_list:
                self.new_participants = "⚠️ Aucun EAN trouvé dans le fichier"
                self.import_enabled = False
                return

            existing_partner_ids = self.community_id.partner_ids.ids
            existing_ean_list = []
            if existing_partner_ids:
                existing_partners = self.env['res.partner'].search(
                    [('id', 'in', existing_partner_ids)])
                existing_ean_list = existing_partners.mapped('ref')

            new_list = []
            existing_list = []
            for ean_data in ean_list:
                ean = ean_data.get('ean', '').strip()
                ptype = ean_data.get('type', 'Prélèvement')
                if not ean:
                    continue
                if ean in existing_ean_list:
                    partner = self.env['res.partner'].search(
                        [('ref', '=', ean)], limit=1)
                    existing_list.append(f"✅ {ean} - {partner.name if partner else ean}")
                else:
                    new_list.append(f"🆕 {ean} - {ptype}")

            self.new_participants = ("\n".join(new_list) if new_list
                                     else "Aucun nouveau participant")
            self.existing_participants = ("\n".join(existing_list) if existing_list
                                          else "Aucun participant existant")
            self.has_new_participants = bool(new_list)

            if (not self.has_new_participants
                    and self.new_participants != "⚠️ Aucun EAN trouvé dans le fichier"):
                self.import_enabled = True
                self.confirm_import = False
            else:
                self.confirm_import = False
                self.import_enabled = False

        except Exception as e:
            self.new_participants = f"⚠️ Erreur de lecture du fichier: {str(e)}"
            self.import_enabled = False
            _logger.error(f"Erreur analyse fichier RESA: {e}")

    def _detect_month(self):
        if not self.file_name:
            return
        patterns = [
            r'(\d{4})[-_ ](\d{1,2})',
            r'(\d{1,2})[-_](\d{4})',
            r'(\d{4})(\d{2})',
            r'(\d{2})[-_](\d{4})',
        ]
        for pattern in patterns:
            match = re.search(pattern, self.file_name)
            if match:
                groups = match.groups()
                if len(groups) == 2:
                    if len(groups[0]) == 4:
                        year, month = groups[0], groups[1]
                    elif len(groups[1]) == 4:
                        month, year = groups[0], groups[1]
                    else:
                        try:
                            y1, m1 = int(groups[0]), int(groups[1])
                            if y1 > 2000 and m1 <= 12:
                                year, month = y1, m1
                            elif m1 > 2000 and y1 <= 12:
                                year, month = m1, y1
                            else:
                                continue
                        except Exception:
                            continue
                    self.month = f"{year}-{int(month):02d}"
                    return

        if self.file_content:
            try:
                excel_data = base64.b64decode(self.file_content)
                workbook = load_workbook(io.BytesIO(excel_data), data_only=True)
                for sheet_name in workbook.sheetnames:
                    sheet = workbook[sheet_name]
                    for row in range(1, min(20, sheet.max_row + 1)):
                        for col in range(1, min(10, sheet.max_column + 1)):
                            val = sheet.cell(row=row, column=col).value
                            if val:
                                val_str = str(val)
                                match = re.search(r'(\d{4})[-/](\d{1,2})', val_str)
                                if match:
                                    year, month = match.groups()
                                    if 2000 <= int(year) <= 2100 and 1 <= int(month) <= 12:
                                        self.month = f"{year}-{int(month):02d}"
                                        return
                                match = re.search(r'(\d{1,2})[-/](\d{4})', val_str)
                                if match:
                                    month, year = match.groups()
                                    if 2000 <= int(year) <= 2100 and 1 <= int(month) <= 12:
                                        self.month = f"{year}-{int(month):02d}"
                                        return
            except Exception as e:
                _logger.warning(f"Erreur détection mois: {e}")

    def _extract_eans_from_sheets(self, workbook):
        ean_list = []
        for sheet_name in workbook.sheetnames:
            if ('partagé' in sheet_name.lower() or 'partage' in sheet_name.lower()
                    or 'shared' in sheet_name.lower() or sheet_name == 'Partagé Rep'):
                sheet = workbook[sheet_name]
                for col in range(2, min(50, sheet.max_column + 1)):
                    ean_val = sheet.cell(row=2, column=col).value
                    if ean_val:
                        ean_str = str(ean_val).strip()
                        digits = re.sub(r'[^0-9]', '', ean_str)
                        ean = digits[:18] if len(digits) >= 14 else ean_str
                        type_val = sheet.cell(row=3, column=col).value
                        ptype = 'Prélèvement'
                        if type_val and 'A-' in str(type_val):
                            ptype = 'Injection'
                        elif type_val and 'A+' in str(type_val):
                            ptype = 'Prélèvement'
                        if len(re.sub(r'[^0-9]', '', ean)) >= 10:
                            ean_list.append({'ean': ean, 'type': ptype})
                break

        if not ean_list:
            for sheet_name in workbook.sheetnames:
                if 'brut' in sheet_name.lower() or sheet_name == 'Brut Rep':
                    sheet = workbook[sheet_name]
                    for col in range(2, min(50, sheet.max_column + 1)):
                        ean_val = sheet.cell(row=2, column=col).value
                        if ean_val:
                            ean_str = str(ean_val).strip()
                            digits = re.sub(r'[^0-9]', '', ean_str)
                            ean = digits[:18] if len(digits) >= 14 else ean_str
                            type_val = sheet.cell(row=3, column=col).value
                            ptype = 'Prélèvement'
                            if type_val and 'A-' in str(type_val):
                                ptype = 'Injection'
                            elif type_val and 'A+' in str(type_val):
                                ptype = 'Prélèvement'
                            if len(re.sub(r'[^0-9]', '', ean)) >= 10:
                                ean_list.append({'ean': ean, 'type': ptype})
                    break
        return ean_list

    @api.onchange('confirm_import')
    def _onchange_confirm_import(self):
        self.import_enabled = self.confirm_import

    # ==========================================
    # ÉTAPE 1
    # ==========================================

    def action_import(self):
        if not self.file_content:
            raise UserError(_("Veuillez sélectionner un fichier."))
        if not self.month or self.month == datetime.now().strftime('%Y-%m'):
            self._detect_month()
        if not self.month:
            self.month = datetime.now().strftime('%Y-%m')

        if self.has_new_participants and not self.confirm_import:
            raise UserError(_("Veuillez confirmer l'ajout des nouveaux participants."))

        return {
            'type': 'ir.actions.act_window',
            'res_model': 'import.resa.wizard',
            'res_id': self.id,
            'view_mode': 'form',
            'view_id': self.env.ref(
                'import_ores_resa.view_import_resa_progress_form').id,
            'target': 'new',
            'name': _('Importation en cours...'),
        }

    # ==========================================
    # ÉTAPE 2 : import AVEC CACHE (correctif)
    # ==========================================

    def action_do_import(self):
        """
        Import synchrone AVEC CACHE.
        Le cache est stocké dans un dictionnaire Python local,
        PAS en tant qu'attribut de l'enregistrement (Odoo interdit).
        """
        self.ensure_one()

        excel_data = base64.b64decode(self.file_content)
        workbook = load_workbook(io.BytesIO(excel_data), data_only=True)

        # 🎯 Cache LOCAL (dictionnaires Python, pas d'attributs d'enregistrement)
        cache = {
            'partners': {},      # {ean: res.partner record}
            'consos': {},        # {(timestamp, ean): conso record}
            'prods': {},         # {(timestamp, ean): prod record}
        }

        # Charger tous les partners en 1 requête
        all_partners = self.env['res.partner'].search([('ref', '!=', False)])
        cache['partners'] = {p.ref: p for p in all_partners}
        _logger.info(f"[RESA] Cache partners: {len(cache['partners'])} entrées")

        # Charger toutes les consommations du mois en 1 requête
        existing_consos = self.env['energy.community.consumption'].search([
            ('community_id', '=', self.community_id.id),
            ('month', '=', self.month),
        ])
        cache['consos'] = {(c.timestamp, c.ean): c for c in existing_consos}
        _logger.info(f"[RESA] Cache consommations: {len(cache['consos'])} entrées")

        # Charger toutes les productions du mois en 1 requête
        existing_prods = self.env['energy.community.production'].search([
            ('community_id', '=', self.community_id.id),
            ('month', '=', self.month),
        ])
        cache['prods'] = {(p.timestamp, p.ean): p for p in existing_prods}
        _logger.info(f"[RESA] Cache productions: {len(cache['prods'])} entrées")

        # Créer les participants manquants
        self._create_participants(workbook, cache)
        self.env.cr.commit()

        # Import des données
        stats = self._import_all_data(workbook, cache)
        self.env.cr.commit()

        self.import_stats = stats

        if self.generate_invoices:
            try:
                self._generate_invoices()
                self.env.cr.commit()
            except Exception as e:
                _logger.warning(f"Erreur génération factures: {e}")

        invoice_stats = (f"\n📊 Factures:\n- Créées: {self.invoices_created}"
                         f"\n- Déjà existantes: {self.invoices_skipped}")
        self.import_stats = stats + invoice_stats

        return {
            'type': 'ir.actions.act_window',
            'res_model': 'energy.community.consumption',
            'view_mode': 'list,form',
            'target': 'main',
            'domain': [('community_id', '=', self.community_id.id),
                       ('month', '=', self.month)],
            'name': f"Données importées - {self.community_id.name} - {self.month}",
            'context': {
                'default_community_id': self.community_id.id,
                'default_month': self.month,
            },
        }

    # ==========================================
    # PARTICIPANTS
    # ==========================================

    def _create_participants(self, workbook, cache):
        ean_list = self._extract_eans_from_sheets(workbook)
        partner_ids_to_add = []
        for ean_data in ean_list:
            ean = ean_data.get('ean', '').strip()
            if not ean:
                continue
            partner = cache['partners'].get(ean)
            if not partner:
                partner = self.env['res.partner'].create({
                    'name': f'Participant {ean}',
                    'ref': ean,
                })
                cache['partners'][ean] = partner
            if partner.id not in self.community_id.partner_ids.ids:
                partner_ids_to_add.append(partner.id)
        if partner_ids_to_add:
            self.community_id.write({
                'partner_ids': [(4, pid) for pid in partner_ids_to_add]
            })

    # ==========================================
    # IMPORT DONNÉES
    # ==========================================

    def _import_all_data(self, workbook, cache):
        stats = {
            'brut': 0, 'partage': 0, 'net': 0, 'errors': 0,
            'start_time': datetime.now(),
        }
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            if 'brut' in sheet_name.lower() or sheet_name == 'Brut Rep':
                stats = self._import_sheet_data(sheet, stats, 'brut', cache)
            elif ('partagé' in sheet_name.lower() or 'partage' in sheet_name.lower()
                  or 'shared' in sheet_name.lower() or sheet_name == 'Partagé Rep'):
                stats = self._import_sheet_data(sheet, stats, 'partage', cache)
            elif 'net' in sheet_name.lower() or sheet_name == 'Net Rep':
                stats = self._import_sheet_data(sheet, stats, 'net', cache)

        stats['end_time'] = datetime.now()
        stats['duration'] = (stats['end_time'] - stats['start_time']).total_seconds()
        return f"""
📊 Statistiques d'import:
- Lignes brutes importées: {stats['brut']}
- Lignes partagées importées: {stats['partage']}
- Lignes nettes importées: {stats['net']}
- Erreurs: {stats['errors']}
- Temps d'exécution: {stats['duration']:.2f} secondes
"""

    def _get_ean_cols_from_sheet(self, sheet):
        ean_cols = {}
        for col in range(2, min(30, sheet.max_column + 1)):
            ean_val = sheet.cell(row=2, column=col).value
            type_val = sheet.cell(row=3, column=col).value
            if ean_val:
                ean_str = str(ean_val).strip()
                digits = re.sub(r'[^0-9]', '', ean_str)
                if len(digits) >= 14:
                    ean = digits[:18]
                else:
                    ean = ean_str
                    if not ean or len(ean) < 10:
                        continue
                ean_type = 'consumption'
                if type_val and 'A-' in str(type_val):
                    ean_type = 'production'
                elif type_val and 'A+' in str(type_val):
                    ean_type = 'consumption'
                if len(re.sub(r'[^0-9]', '', ean)) >= 10:
                    ean_cols[col] = {'ean': ean, 'type': ean_type}
        return ean_cols

    def _parse_timestamp(self, value):
        if not value:
            return None
        if isinstance(value, datetime):
            return value
        if isinstance(value, str):
            value_str = value.strip()
            formats = [
                '%Y-%m-%dT%H:%M:%S+%z', '%Y-%m-%dT%H:%M:%S',
                '%Y-%m-%d %H:%M:%S+%z', '%Y-%m-%d %H:%M:%S', '%Y-%m-%d %H:%M',
            ]
            for fmt in formats:
                try:
                    return datetime.strptime(value_str, fmt)
                except ValueError:
                    continue
            match = re.search(r'(\d{4}-\d{2}-\d{2})[T ](\d{2}:\d{2}:\d{2})', value_str)
            if match:
                try:
                    return datetime.strptime(
                        f"{match.group(1)} {match.group(2)}", '%Y-%m-%d %H:%M:%S')
                except ValueError:
                    pass
        if isinstance(value, (int, float)) and 40000 < value < 50000:
            try:
                return datetime.fromordinal(int(value) - 693594)
            except Exception:
                pass
        return None

    def _get_data_start_row(self, sheet):
        for row in range(5, min(100, sheet.max_row + 1)):
            val = sheet.cell(row=row, column=1).value
            if val and self._parse_timestamp(val):
                return row
        return 5

    def _import_sheet_data(self, sheet, stats, sheet_type, cache):
        """
        Importe une feuille avec cache.
        """
        ean_cols = self._get_ean_cols_from_sheet(sheet)
        if not ean_cols:
            stats['errors'] += 1
            return stats
        data_start_row = self._get_data_start_row(sheet)
        if not data_start_row:
            stats['errors'] += 1
            return stats

        imported_count = 0

        for row in range(data_start_row, sheet.max_row + 1):
            timestamp = None
            val = sheet.cell(row=row, column=1).value
            if val:
                timestamp = self._parse_timestamp(val)
            if not timestamp:
                continue

            first_val = sheet.cell(row=row, column=1).value
            if first_val and isinstance(first_val, (int, float)) and first_val in [0, 1, 2]:
                continue
            if (first_val and isinstance(first_val, str)
                    and first_val in ['0', 'Tous', 'Total', 'Totaux']):
                continue

            for col, info in ean_cols.items():
                value = sheet.cell(row=row, column=col).value
                if value and isinstance(value, (int, float)) and abs(value) > 0.0001:
                    try:
                        if info['type'] == 'consumption':
                            self._update_consumption_cached(
                                timestamp, info['ean'], value, sheet_type, cache)
                        elif info['type'] == 'production':
                            self._update_production_cached(
                                timestamp, info['ean'], value, sheet_type, cache)
                        imported_count += 1
                    except Exception as e:
                        _logger.warning(
                            f"Erreur import cellule (row={row}, col={col}): {e}")
                        stats['errors'] += 1

            if imported_count > 0 and imported_count % BATCH_SIZE == 0:
                self.env.cr.commit()
                _logger.info(f"[RESA-{sheet_type}] Commit après {imported_count} lignes")

        self.env.cr.commit()

        if sheet_type == 'brut':
            stats['brut'] = imported_count
        elif sheet_type == 'partage':
            stats['partage'] = imported_count
        elif sheet_type == 'net':
            stats['net'] = imported_count
        return stats

    # ==========================================
    # MISE À JOUR VIA CACHE
    # ==========================================

    def _update_consumption_cached(self, timestamp, ean, value, sheet_type, cache):
        key = (timestamp, ean)
        record = cache['consos'].get(key)

        if not record:
            partner = cache['partners'].get(ean)
            if not partner:
                partner = self.env['res.partner'].create({
                    'name': f'Participant {ean}',
                    'ref': ean,
                })
                cache['partners'][ean] = partner

            record = self.env['energy.community.consumption'].create({
                'timestamp': timestamp,
                'ean': ean,
                'partner_id': partner.id,
                'community_id': self.community_id.id,
                'month': self.month,
                'iteration': 0,
            })
            cache['consos'][key] = record

        if sheet_type == 'brut':
            record.prelevement_brut = value
        elif sheet_type == 'partage':
            record.prelevement_couvert = value
        elif sheet_type == 'net':
            record.allo_consommation = value

    def _update_production_cached(self, timestamp, ean, value, sheet_type, cache):
        key = (timestamp, ean)
        record = cache['prods'].get(key)

        if not record:
            partner = cache['partners'].get(ean)
            if not partner:
                partner = self.env['res.partner'].create({
                    'name': f'Participant {ean}',
                    'ref': ean,
                })
                cache['partners'][ean] = partner

            record = self.env['energy.community.production'].create({
                'timestamp': timestamp,
                'ean': ean,
                'partner_id': partner.id,
                'community_id': self.community_id.id,
                'month': self.month,
            })
            cache['prods'][key] = record

        if sheet_type == 'brut':
            record.production_brute = value
        elif sheet_type == 'partage':
            record.production_autoconsommee = value
        elif sheet_type == 'net':
            record.production_non_allouee = value

    # ==========================================
    # FACTURATION
    # ==========================================

    def _get_tax_6_percent(self):
     tax_group = self.env['account.tax.group'].search([('name', 'ilike', '6%')], limit=1)
     tax = self.env['account.tax'].search([
        ('name', '=', 'TVA 6%'),
        ('company_id', '=', self.env.company.id)], limit=1)
     if not tax:
        tax = self.env['account.tax'].create({
            'name': 'TVA 6%', 'amount': 6.0, 'amount_type': 'percent',
            'type_tax_use': 'sale', 'tax_exigibility': 'on_invoice',
            'description': 'TVA 6% (énergie)', 'company_id': self.env.company.id,
            'tax_group_id': tax_group.id if tax_group else False,
        })
     return tax

    def _get_invoice_journal(self, file_type):
        journal_type = 'purchase' if file_type == 'production' else 'sale'
        return self.env['account.journal'].search([('type', '=', journal_type)], limit=1)

    def _get_payment_term_30_days(self):
        payment_term = self.env['account.payment.term'].search(
            [('name', 'ilike', '30%')], limit=1)
        if not payment_term:
            payment_term = self.env['account.payment.term'].create({
                'name': '30 jours fin de mois',
                'note': 'Paiement à 30 jours fin de mois',
                'line_ids': [(0, 0, {
                    'value': 'balance', 'days': 30,
                    'days_type': 'end_of_month', 'sequence': 1,
                })]
            })
        return payment_term

    def _invoice_exists(self, partner_id, move_type, month, community_id, ean):
        ref_prefix = f'Opération de partage {community_id.name} - {month} - {ean}'
        existing = self.env['account.move'].search([
            ('partner_id', '=', partner_id),
            ('move_type', '=', move_type),
            ('state', 'in', ['draft', 'posted']),
            ('ref', 'ilike', ref_prefix),
        ], limit=1)
        return bool(existing)

    def _generate_invoices(self):
        self.invoices_created = 0
        self.invoices_skipped = 0
        tax_6 = self._get_tax_6_percent()
        payment_term = self.payment_term_id or self._get_payment_term_30_days()
        self._generate_consumption_invoices(tax_6, payment_term)
        self._generate_production_invoices(tax_6, payment_term)

    def _generate_consumption_invoices(self, tax_6, payment_term):
        journal = self._get_invoice_journal('consumption')
        if not journal:
            return
        consumptions = self.env['energy.community.consumption'].search([
            ('community_id', '=', self.community_id.id),
            ('month', '=', self.month)])
        consumption_by_ean = {}
        for cons in consumptions:
            ean = cons.ean
            if ean not in consumption_by_ean:
                consumption_by_ean[ean] = {
                    'prelevement_couvert': 0.0,
                    'partner_id': cons.partner_id.id}
            consumption_by_ean[ean]['prelevement_couvert'] += cons.prelevement_couvert

        product_price = self.community_id.product_price or 0.0
        if not self.env.user.has_group('account.group_account_invoice'):
            return

        for ean, data in consumption_by_ean.items():
            if data['prelevement_couvert'] > 0 and product_price > 0:
                partner = self.env['res.partner'].browse(data['partner_id'])
                if self._invoice_exists(partner.id, 'out_invoice',
                                        self.month, self.community_id, ean):
                    self.invoices_skipped += 1
                    continue
                quantity = data['prelevement_couvert']
                unit_price = product_price
                total = quantity * unit_price
                description = (
                    f"Consommation énergie partagée - {self.month}\n"
                    f"Prélèvement couvert : {quantity:.3f} kWh\n"
                    f"Prix de vente : {unit_price:.3f} €/kWh\n"
                    f"Sous-total : {total:.2f} €")
                invoice_lines = [{
                    'name': description, 'quantity': quantity,
                    'price_unit': unit_price, 'tax_ids': [(6, 0, [tax_6.id])]}]

                # === NOUVELLES LIGNES : CV + ACCISES (facture CLIENT uniquement) ===
                cv_total = quantity * CV_RESTITUTION_PRICE
                invoice_lines.append({
                    'name': (
                        f"Coût de la restitution des CV - {self.month}\n"
                        f"Prélèvement couvert : {quantity:.3f} kWh\n"
                        f"Tarif : {CV_RESTITUTION_PRICE:.3f} €/kWh HTVA\n"
                        f"Sous-total : {cv_total:.2f} €"),
                    'quantity': quantity,
                    'price_unit': CV_RESTITUTION_PRICE,
                    'tax_ids': [(6, 0, [tax_6.id])],
                })
                accises_total = quantity * ACCISES_2026_PRICE
                invoice_lines.append({
                    'name': (
                        f"Droit d'accises spécial 2026 (tarif unique) - {self.month}\n"
                        f"Prélèvement couvert : {quantity:.3f} kWh\n"
                        f"Tarif : {ACCISES_2026_PRICE:.3f} €/kWh HTVA\n"
                        f"Sous-total : {accises_total:.2f} €"),
                    'quantity': quantity,
                    'price_unit': ACCISES_2026_PRICE,
                    'tax_ids': [(6, 0, [tax_6.id])],
                })

                self._create_draft_invoice(partner, invoice_lines, ean,
                                           journal, payment_term, 'out_invoice')
                self.invoices_created += 1

    def _generate_production_invoices(self, tax_6, payment_term):
        journal = self._get_invoice_journal('production')
        if not journal:
            return
        productions = self.env['energy.community.production'].search([
            ('community_id', '=', self.community_id.id),
            ('month', '=', self.month)])
        production_by_ean = {}
        for prod in productions:
            ean = prod.ean
            if ean not in production_by_ean:
                production_by_ean[ean] = {
                    'production_autoconsommee': 0.0,
                    'partner_id': prod.partner_id.id}
            production_by_ean[ean]['production_autoconsommee'] += prod.production_autoconsommee

        producer_price = self.community_id.producer_price or 0.0
        if not self.env.user.has_group('account.group_account_invoice'):
            return

        for ean, data in production_by_ean.items():
            if data['production_autoconsommee'] > 0 and producer_price > 0:
                partner = self.env['res.partner'].browse(data['partner_id'])
                if self._invoice_exists(partner.id, 'in_invoice',
                                        self.month, self.community_id, ean):
                    self.invoices_skipped += 1
                    continue
                quantity = data['production_autoconsommee']
                unit_price = producer_price
                total = quantity * unit_price
                description = (
                    f"Rémunération énergie partagée - {self.month}\n"
                    f"Production autoconsommée : {quantity:.3f} kWh\n"
                    f"Prix d'achat : {unit_price:.3f} €/kWh\n"
                    f"Sous-total : {total:.2f} €")
                invoice_lines = [{
                    'name': description, 'quantity': quantity,
                    'price_unit': unit_price, 'tax_ids': [(6, 0, [tax_6.id])]}]
                self._create_draft_invoice(partner, invoice_lines, ean,
                                           journal, payment_term, 'in_invoice')
                self.invoices_created += 1

    def _create_draft_invoice(self, partner, invoice_lines, ean, journal,
                              payment_term, move_type):
        try:
            line_vals_list = []
            for line in invoice_lines:
                line_vals = {
                    'name': line['name'], 'quantity': line['quantity'],
                    'price_unit': line['price_unit'],
                    'tax_ids': line.get('tax_ids', [(6, 0, [])])}
                line_vals_list.append((0, 0, line_vals))
            invoice_vals = {
                'partner_id': partner.id, 'move_type': move_type,
                'invoice_date': self.invoice_date or fields.Date.today(),
                'journal_id': journal.id,
                'ref': f'Opération de partage {self.community_id.name} - {self.month} - {ean}',
                'invoice_line_ids': line_vals_list,
                'invoice_payment_term_id': payment_term.id,
                'state': 'draft'}
            self.env['account.move'].create(invoice_vals)
        except Exception as e:
            _logger.warning(f"Erreur création facture pour {ean}: {e}")